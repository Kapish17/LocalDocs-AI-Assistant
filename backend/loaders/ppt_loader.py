from langchain_core.documents import Document
from pptx import Presentation


def load_ppt(file_path):
    """
    Reads a PowerPoint file and extracts text, one Document per slide.
    Slides map naturally to "page" for citation purposes.
    """

    presentation = Presentation(file_path)

    slides = []
    for slide_num, slide in enumerate(presentation.slides, start=1):
        text = ""
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
        slides.append(
            Document(
                page_content=text,
                metadata={
                    "source": file_path.name,
                    "file_type": "pptx",
                    "document_id": file_path.name,
                    "page": slide_num,
                },
            )
        )

    return slides
